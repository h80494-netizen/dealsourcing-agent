"""
VC 투자심사역(Investment Committee) 관점 기업분석 프롬프트 자동화 및 리포트 생성기
- 42개 업종 분류 체계(대·중·소분류 및 기존/신종 산업)
- 6대 핵심 투자심사 항목 (Cap Table, 시장, 기업/경영진, 프라이싱, 자금용도, 자본조달)
- 5대 섹터별(AI, ICT, 바이오, 커머스, 기타) 특화 지침 동적 결합
"""

import os
import io
import json
import logging
from typing import Optional, Dict, Any, Tuple

logger = logging.getLogger(__name__)

# 업종별 특화 심사 지침 모듈
INDUSTRY_SPECIFIC_MODULES = {
    "AI": """[AI (인공지능) 업종 추가 심사 지시사항]
- 기존 산업 (AI-T-01~02: 예측형 AI·비전·NLP): SI성 용역 매출 비중과 반복 매출(ARR) 비중을 엄격히 구분하십시오. 인건비 투입 대비 매출 총이익률(Gross Margin)이 60% 이상 나오는 제품화(Productization)가 완료되었는지 검증하십시오.
- 신종 산업 (AI-E-01~03: 생성형·에이전틱·피지컬 AI): 빅테크(OpenAI, Google 등)의 모델 업데이트 시 사업 모델이 소멸할 위험(Wrapper Risk)을 평가하십시오. 피지컬 AI의 경우 시뮬레이션(Sim-to-Real) 데이터 확보 역량과 하드웨어(액추에이터/센서) 통합 제어 능력을 집중 분석하십시오.
- 신종 산업 (AI-E-04~05: AI 반도체·인프라): 자금용도에서 파운드리 공정(Tape-out) 비용과 클라우드 GPU 서버 임차 비용(Compute Burn Rate)을 면밀히 따지십시오. 프라이싱 시 단순 기술 기대감이 아닌 실제 고객사 PoC 통과 및 양산 수주(Design Win) 여부를 반영하십시오.""",

    "ICT": """[ICT (정보통신·SW·HW) 업종 추가 심사 지시사항]
- 기존 산업 (IC-T-01~03: 엔터프라이즈 SW·모바일·통신HW): 시장 포화도와 고객 이탈률(Churn Rate)을 최우선으로 보십시오. 자본조달 시 과도한 밸류에이션의 보통주/RCPS보다는 현금흐름에 기반한 메자닌(CB/BW)이나 배당 조건부 투자가 적합한지 검토하십시오.
- 신종 산업 (IC-E-01~02: 버티컬 SaaS·차세대 보안): 순매출유지율(NRR 110% 이상 여부), LTV/CAC 비율(3배 이상 여부), CAC 회수기간(Payback Period)을 기준으로 프라이싱 멀티플(PSR)의 적정성을 검증하십시오.
- 신종 산업 (IC-E-03~05: 양자기술·공간컴퓨팅·웹3): 상용화까지 긴 시간이 걸리는 딥테크이므로 정부 R&D 과제(무상 자금) 조달 능력과 핵심 특허 포트폴리오를 체크하십시오. 웹3의 경우 토큰 이코노미와 지분(Equity) 가치가 충돌하지 않는지(가치 귀속 구조) 법적·구조적으로 분석하십시오.""",

    "바이오": """[바이오 (바이오·헬스케어) 업종 추가 심사 지시사항]
- 기존 산업 (BI-T-01~03: 제네릭·일반 의료기기·건기식·CMO): 공장 증설(GMP 인증) 등 시설투자(CAPEX) 비중이 높으므로, 가동률(Utilization Rate)과 수주 잔고에 기반해 유상증자 외에 시설자금 차입이나 CB/BW 활용 가능성을 분석하십시오.
- 신종 산업 (BI-E-01~02: 차세대 모달리티 CGT/ADC·합성생물학): 임상 단계(비임상/1상/2상)별 성공 확률을 반영한 위험조정 순현재가치(rNPV)와 글로벌 빅파마의 최근 기술이전(L/O) 계약 규모(Upfront 비율 및 마일스톤)를 기준으로 프라이싱을 검증하십시오. 자금용도가 임상시험수탁기관(CRO) 비용과 CDMO 시료 생산에 정확히 매칭되는지 확인하십시오.
- 신종 산업 (BI-E-03~05: AI 신약·디지털치료제·뇌과학/역노화): FDA/식약처 인허가 통과 여부뿐 아니라, 건강보험 수가(Reimbursement) 적용 가능성과 실제 일선 의사들의 처방 유인책을 시장 진입장벽의 핵심으로 분석하십시오.""",

    "커머스": """[커머스 (유통·소비재) 업종 추가 심사 지시사항]
- 기존 산업 (CM-T-01~03: 종합몰·F&B/패션·로컬 O2O): 거래액(GMV)의 착시를 제거하고 실질 수수료율(Take Rate)과 공헌이익(Contribution Margin) 흑자 여부를 검증하십시오. 재고자산 회전율과 운전자본(Working Capital) 묶임 현상을 분석해 자본조달 수단(단기 차입 vs 지분 투자)을 구분하십시오.
- 신종 산업 (CM-E-01~04: 크로스보더·리커머스·리테일테크·슬립/펨/펫테크): 특정 플랫폼(아마존, 틱톡샵 등) 알고리즘 변경 리스크, 해외 물류/관세 장벽, 재구매율(Retention Cohort)을 집중 분석하십시오. 자금용도에서 단순 퍼포먼스 마케팅(광고비 태우기) 비중이 과도한 경우 프라이싱을 보수적으로 깎으십시오.""",

    "기타": """[기타 (기후·모빌리티·로봇·핀테크·콘텐츠) 업종 추가 심사 지시사항]
- 기존 산업 (OT-T-01~04: 전통 소부장·물류·초기 핀테크·게임/미디어): 전방 산업(완성차, 반도체, 건설 등)의 경기 사이클 의존도와 특정 대기업 매출 편중 리스크(단가 인하 압력)를 분석하십시오. 게임/미디어는 단일 프로젝트 흥행 리스크를 줄일 수 있는 파이프라인 다각화 여부를 체크하십시오.
- 신종 산업 (OT-E-01~02: 기후테크·클린에너지): 정부 보조금이나 탄소배출권 가격 변동 없이도 자체적인 경제성(Unit Economics, 그린 프리미엄 극복 여부)이 나오는 시점을 산출하십시오. 대규모 파일럿 플랜트 설비투자가 필요하므로 프로젝트 파이낸싱(PF), 정책자금 차입, 메자닌(CB/BW)을 결합한 자본조달 구조를 설계하십시오.
- 신종 산업 (OT-E-03~05: 미래 모빌리티·우주/방산·로보틱스): 방산/우주/로봇은 수주 리드타임이 길고 양산 전환 시점에 막대한 운전자본이 소요되는 '죽음의 계곡(Valley of Death)'이 존재합니다. 이번 펀딩 규모가 양산 라인 구축(CAPEX)까지 버틸 수 있는지, 창업자가 기술자 마인드를 넘어 제조·양산·품질관리(QA) 역량을 갖추었는지 검증하십시오.
- 신종 산업 (OT-E-06~08: B2B 핀테크·애그/푸드테크·크리에이터/IP): 금융 규제(라이선스 진입장벽), 식약처 신소재 승인(배양육 등), 핵심 크리에이터/IP 이탈 리스크(Key-man Risk)를 집중 점검하고, 주주간계약서(SHA)에 핵심 인력 의무재직(Lock-up) 및 IP 귀속 조항을 반드시 포함하도록 제안하십시오."""
}

MASTER_IC_TEMPLATE = """# Role: 글로벌 Top-Tier 벤처캐피탈(VC) 수석 투자심사역 (Partner & IC Member)

# Task
제공된 기업의 데이터(IR 자료, 재무제표, 뉴스, 사업계획서, 다국어 기업 정보 등)를 바탕으로, 우리 하우스의 표준 산업 분류 체계에 따라 업종을 매핑하고 아래 6가지 핵심 심사 기준 및 업종별 특화 기준에 맞춰 냉철하고 입체적인 투자심사보고서(Investment Committee Memo)를 한국어로 작성하십시오.

---

## [Step 0] 산업 분류 및 포지셔닝 매핑
제공된 기업 설명과 키워드(한국어/영어/중국어/일본어/독일어/프랑스어)를 분석하여 아래 항목을 먼저 정의하십시오.
- **분류 코드 및 경로:** {industry_mapping}
- **포지셔닝 진단:** 해당 기업이 '기존 산업(안정적 현금흐름·시장점유율 싸움)'인지 '신종 산업(선점효과·기술 해자·밸류에이션 리레이팅)'인지 정의하고, 이에 맞는 심사 잣대를 적용하십시오.

---

## [Step 1] 6대 핵심 투자심사 분석 (Deep-Dive Analysis)

### 1. 지분구성 및 펀드레이징 스토리 (Cap Table & Fundraising Equity Story)
- **주주 구성(Cap Table) 건전성:** 창업자(대표이사) 및 핵심 경영진의 현재 지분율과 이번 라운드 및 향후 Exit(IPO/M&A) 시점까지 예상되는 지분 희석(Dilution) 후 경영권 방어 가능성을 평가하십시오. (Dead Equity: 퇴사한 공동창업자나 초기 비활동 투자자 지분 유무 체크)
- **펀드레이징 히스토리:** 시드(Seed)부터 현재 라운드까지의 투자 유치 시점, 밸류에이션 추이(Step-up 배수), 기존 투자사(FI/SI)의 면면 및 후속 투자(Follow-on) 참여 여부를 분석하십시오.
- **스토리의 일관성:** 이전 라운드에서 투자자들에게 약속했던 마일스톤(KPI)을 실제로 달성했는지, 피벗(Pivot)이 있었다면 그 논리적 타당성이 충분한지 검증하십시오.

### 2. 시장 분석 (Market Dynamics: 성장성·확장성·규모·진입장벽)
- **시장 규모 (TAM-SAM-SOM):** 단순 글로벌 통계 인용이 아닌, 해당 기업의 제품 단가(ARPU)와 타깃 고객 수를 기반으로 한 Bottom-up 방식의 실질 유효 시장(SOM) 규모를 산출하십시오.
- **성장성(CAGR) 및 타이밍(Why Now?):** 시장의 연평균 성장률과 함께, 왜 하필 지금 이 시장이 폭발하는지(규제 변화, 기술 단가 하락, 소비자 행동 변화 등)를 분석하십시오.
- **글로벌 및 인접 시장 확장성(Scalability):** 한계비용 증가 없이 타 지역(크로스보더)이나 인접 버티컬 산업으로 확장할 수 있는 구조인지 평가하십시오.
- **구조적 시장 진입장벽:** 인허가/규제, 초기 자본 집약도, 네트워크 효과, 전환 비용(Switching Cost) 등 후발주자의 진입을 막는 시장 자체의 장벽을 분석하십시오.

### 3. 기업 및 경영진 분석 (Company & Founder: 기술력·해자·오너 역량)
- **기술력 및 독점적 해자(Moat):** 경쟁사 대비 기술적 우위(특허, 독점 데이터셋, 수율, 원가 경쟁력, 알고리즘 성능 등)와 경쟁자가 이를 복제하는 데 걸리는 시간(Lead Time)을 평가하십시오.
- **오너(창업자) 및 핵심 팀 역량 (Founder-Market Fit):** 대표이사의 도메인 전문성, 과거 성공/실패 경험, 기술/영업 핵심 인력(C-level)의 구성, 그리고 위기 상황에서의 실행력과 자본 배분(Capital Allocation) 능력을 냉정하게 평가하십시오.

### 4. 프라이싱 및 밸류에이션 (Pricing: 가격 적정성·펀딩 규모)
- **Pre/Post-money 밸류에이션 적정성:** 동종 업계(Peer Group) 상장사 및 최근 비상장 M&A/투자 거래의 멀티플(PSR, PER, EV/EBITDA, 파이프라인 가치 등)과 비교하여 현재 요구하는 기업가치가 고평가(Overvalued)인지 할인(Discount) 상태인지 산출하십시오.
- **펀딩 규모의 적정성:** 현재 회사의 월간 현금소진율(Burn Rate)과 런웨이(Runway)를 고려할 때, 이번 목표 펀딩 금액이 과도한 지분 희석 없이 다음 라운드(또는 BEP 달성)의 마일스톤을 달성하기에 충분한 금액인지 검증하십시오.
- **예상 투자수익률 (Exit Scenario):** 보수적·기본·낙관적 시나리오별 예상 Exit 시점, Exit 기업가치, 목표 IRR 및 멀티플(MOIC)을 시뮬레이션하십시오.

### 5. 자금 용도 (Use of Proceeds: 인력·시설·R&D·운전자본)
- **자금 투입 계획의 구체성:** 조달 자금이 ① 핵심 인력 충원(R&D/영업), ② 시설 및 장비 투자(CAPEX/인프라), ③ 연구개발 및 임상/인증, ④ 마케팅 및 운전자본(OPEX)에 각각 몇 %씩 배분되는지 분석하십시오.
- **자본 효율성(Capital Efficiency) 검증:** 조달된 자금이 단순 과거 부채 상환이나 비효율적 고정비 유지에 쓰이지 않고, 직접적인 기업가치 상승(Value-up)과 매출/기술 마일스톤 달성에 직결되는지 평가하십시오.

### 6. 자본조달 방법 및 구조 (Financing Structure: 증자·CB·BW·차입 등)
- **최적 조달 수단 검토:** 기업의 현금흐름, 담보력, 리스크 수준을 고려할 때 아래 수단 중 무엇이 투자자와 기업 양쪽에 가장 유리한지 제안하십시오.
  1) **보통주/상환전환우선주(RCPS)/전환우선주(CPS) 유상증자:** 리픽싱(Refixing) 조건, 상환권, 우선청산권(Liquidation Preference) 설계
  2) **메자닌 (CB 전환사채 / BW 신주인수권부사채):** 하방 경직성 확보(원리금 보장 + 전환권/워런트 업사이드), 만기보장수익률(YTM) 및 표면이자율, 콜옵션(Call Option)/풋옵션(Put Option) 비율
  3) **벤처대출(Venture Debt) 및 일반 차입:** 지분 희석을 최소화하기 위한 에쿼티+데트(Equity+Debt) 혼합 조달 가능성
- **딜 구조(Term Sheet) 핵심 특약 제안:** 투자자 보호를 위해 반드시 삽입해야 할 주요 계약 조건(Tag-along, Drag-along, 주요 경영사항 동의권, 경업금지, 마일스톤 분할 납입 등)을 제시하십시오.

---

## [Step 2] 업종 특화 핵심 체크포인트
{industry_specific_instructions}

---

## [Step 3] 최종 투자 의견 (Investment Verdict)
- **투자 매력도 종합 등급:** 강력 추천(A) / 긍정적 검토(B) / 조건부 보류(C) / 투자 부적합(D)
- **핵심 투자 포인트 (Key Investment Highlights):** 3가지 요약
- **핵심 리스크 및 완화 방안 (Key Risks & Mitigation):** 3가지 요약
- **IC 심의 권고 코멘트:** 파트너/심사위원 관점의 결론 제언

---

# 입력 데이터 (기업 정보 및 제출 자료):
- **분석 대상 기업명:** {company_name}
- **산업 분류 기준:** {industry_mapping}
- **기업 제출 데이터 (IR, 재무제표, 사업내용 등):**
{materials_content}
"""

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """PDF 바이트 데이터에서 텍스트를 추출"""
    extracted_text = []
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                extracted_text.append(f"--- [Page {page_idx + 1}] ---\n" + text.strip())
        return "\n\n".join(extracted_text)
    except ImportError:
        try:
            import pypdf2
            reader = pypdf2.PdfReader(io.BytesIO(file_bytes))
            for page_idx, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    extracted_text.append(f"--- [Page {page_idx + 1}] ---\n" + text.strip())
            return "\n\n".join(extracted_text)
        except Exception as e:
            logger.warning(f"PDF extraction fallback error: {e}")
            return "[PDF 텍스트 추출 라이브러리 미설치 또는 파싱 실패]"
    except Exception as e:
        logger.error(f"Error extracting text from PDF: {e}")
        return f"[PDF 텍스트 추출 중 오류 발생: {str(e)}]"


def extract_text_from_docx(file_bytes: bytes) -> str:
    """MS Word (.docx) 파일에서 텍스트 및 표 데이터를 추출"""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        parts = []

        # 1. 단락(Paragraphs) 추출
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                parts.append(text)

        # 2. 표(Tables) 데이터 추출
        if doc.tables:
            parts.append("\n[문서 내 표(Table) 데이터 요약]")
            for table_idx, table in enumerate(doc.tables):
                table_lines = []
                for row in table.rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    # 중복 셀(merged cells) 정리
                    dedup_cells = []
                    for c in row_cells:
                        if not dedup_cells or dedup_cells[-1] != c:
                            dedup_cells.append(c)
                    if any(dedup_cells):
                        table_lines.append(" | ".join(dedup_cells))
                if table_lines:
                    parts.append(f"--- 표 {table_idx + 1} ---\n" + "\n".join(table_lines))

        return "\n\n".join(parts)
    except Exception as e:
        logger.error(f"Error extracting text from Word docx: {e}")
        return f"[Word(.docx) 파일 텍스트 추출 중 오류: {str(e)}]"


def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
    """확장자별(PDF, DOCX, TXT 등) 통합 텍스트 추출 함수"""
    ext = os.path.splitext(filename)[1].lower() if filename else ""

    if ext == ".pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext in [".docx"]:
        return extract_text_from_docx(file_bytes)
    elif ext in [".txt", ".csv", ".json", ".md"]:
        for enc in ['utf-8', 'cp949', 'euc-kr', 'latin-1']:
            try:
                return file_bytes.decode(enc)
            except Exception:
                continue
        return file_bytes.decode('utf-8', errors='ignore')
    elif ext == ".doc":
        # 구형 .doc 바이너리 파일: 문자열 검색 추출 시도
        try:
            raw_str = file_bytes.decode('latin-1', errors='ignore')
            import re
            words = re.findall(r'[\w가-힣\s.,!?:;%()/\-]{4,}', raw_str)
            return "\n".join(words[:2000])
        except Exception as e:
            return f"[구형 .doc 파일 파싱 제한 (최신 .docx 파일 저장을 권장합니다): {e}]"
    else:
        try:
            return file_bytes.decode('utf-8', errors='ignore')
        except Exception:
            return f"[{filename}: 지원하지 않는 파일 형식이거나 읽을 수 없습니다]"


def scrape_website_content(url: str, timeout: int = 8) -> str:
    """기업 홈페이지 URL에서 메타데이터 및 주요 텍스트를 크롤링하여 정제"""
    import urllib.request
    import urllib.parse
    import re
    from bs4 import BeautifulSoup

    if not url.startswith("http://") and not url.startswith("https://"):
        url = "https://" + url

    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7"
    }

    try:
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, timeout=timeout) as response:
            html = response.read().decode('utf-8', errors='ignore')

        soup = BeautifulSoup(html, 'html.parser')

        # 불필요한 태그 제거
        for tag in soup(['script', 'style', 'noscript', 'svg', 'header', 'footer', 'nav', 'iframe']):
            tag.decompose()

        title = soup.title.string.strip() if soup.title and soup.title.string else ""
        meta_desc = ""
        desc_tag = soup.find('meta', attrs={'name': 'description'}) or soup.find('meta', attrs={'property': 'og:description'})
        if desc_tag and desc_tag.get('content'):
            meta_desc = desc_tag['content'].strip()

        # 주요 텍스트 추출 (h1~h4, p, li)
        text_elements = []
        for elem in soup.find_all(['h1', 'h2', 'h3', 'h4', 'p', 'li']):
            t = elem.get_text(strip=True)
            if len(t) >= 15:  # 너무 짧은 메뉴 텍스트 제외
                text_elements.append(t)

        combined_text = "\n".join(text_elements)
        # 중복 라인 정리
        lines = []
        seen = set()
        for line in combined_text.split('\n'):
            line_str = line.strip()
            if line_str and line_str not in seen:
                seen.add(line_str)
                lines.append(line_str)

        result_summary = f"[기업 공식 웹사이트 실시간 크롤링 정보: {url}]\n"
        if title:
            result_summary += f"- 사이트 제목: {title}\n"
        if meta_desc:
            result_summary += f"- 사이트 설명: {meta_desc}\n"
        result_summary += "- 주요 소개 및 본문 내용:\n" + "\n".join(lines[:120])  # 최대 120개 핵심 문장

        return result_summary
    except Exception as e:
        logger.warning(f"Failed to scrape website {url}: {e}")
        return f"[기업 웹사이트 크롤링 오류 ({url}): {str(e)}]"


def extract_urls_from_text(text: str) -> list:
    """텍스트 내 포함된 웹사이트 URL 목록 추출"""
    import re
    # http://, https:// 및 www. 형태의 URL 추출
    pattern = r'(https?://[a-zA-Z0-9.\-_~:/?#\[\]@!$&\'()*+,;=%]+|www\.[a-zA-Z0-9.\-_~:/?#\[\]@!$&\'()*+,;=%]+)'
    matches = re.findall(pattern, text)
    cleaned_urls = []
    for m in matches:
        u = m.rstrip('.,!?:;')
        if len(u) > 5 and '.' in u:
            cleaned_urls.append(u)
    return list(dict.fromkeys(cleaned_urls))


def build_ic_memo_prompt(
    company_name: str,
    industry_info: Dict[str, Any],
    materials_content: str
) -> Tuple[str, str]:
    """
    산업분류 및 기업자료를 결합한 마스터 IC Memo 프롬프트 생성
    Returns:
        (prompt_text, major_category)
    """
    major = industry_info.get("major", "기타")
    middle = industry_info.get("middle", "기타/미정")
    minor = industry_info.get("minor", "기타")
    code = industry_info.get("code", "OT-T-01")
    ind_type = industry_info.get("type", "신종 산업" if "-E-" in code else "기존 산업")

    mapping_str = f"{code} | {major} > {ind_type} > {middle} > {minor}"
    specific_inst = INDUSTRY_SPECIFIC_MODULES.get(major, INDUSTRY_SPECIFIC_MODULES["기타"])

    # 텍스트가 너무 길면 LLM 컨텍스트 한도 고려 (최대 100,000자)
    if len(materials_content) > 100000:
        materials_content = materials_content[:100000] + "\n\n[...이하 내용 분량 초과로 축약됨...]"

    final_prompt = MASTER_IC_TEMPLATE.format(
        company_name=company_name,
        industry_mapping=mapping_str,
        industry_specific_instructions=specific_inst,
        materials_content=materials_content if materials_content.strip() else "[제공된 상세 자료 없음 - 기업명 및 공개 정보 기반 추정 분석 실시]"
    )

    return final_prompt, major


async def generate_ic_memo(
    company_name: str,
    industry_info: Dict[str, Any],
    materials_content: str,
    api_key: Optional[str] = None
) -> str:
    """
    Gemini 모델을 호출하여 심층 IC Memo 리포트 생성
    """
    prompt, major = build_ic_memo_prompt(company_name, industry_info, materials_content)

    # API Key 확인
    key = api_key or os.getenv("GEMINI_API_KEY")
    if not key:
        return (
            "### [오류] Gemini API 키가 설정되지 않았습니다.\n"
            "환경변수 `GEMINI_API_KEY`를 설정하거나 관리자에게 문의하세요.\n\n"
            "#### 생성된 프롬프트 스크립트 (수동 LLM 복사용):\n\n"
            "```markdown\n" + prompt + "\n```"
        )

    try:
        import google.generativeai as genai
        genai.configure(api_key=key)

        model_candidates = [
            "gemini-flash-lite-latest",
            "gemini-2.0-flash",
            "gemini-3.8-flash",
            "gemini-2.0-flash-exp",
            "gemini-pro"
        ]

        last_error = None
        for model_name in model_candidates:
            try:
                model = genai.GenerativeModel(model_name)
                response = model.generate_content(prompt)
                if response and response.text:
                    return response.text
            except Exception as e:
                last_error = e
                logger.warning(f"Model {model_name} failed: {e}, trying next...")

        return (
            f"### [오류] 리포트 생성 실패\n"
            f"AI 모델 응답 중 오류가 발생했습니다: {last_error}\n\n"
            "#### 생성된 프롬프트 스크립트:\n\n"
            "```markdown\n" + prompt + "\n```"
        )

    except Exception as e:
        logger.error(f"Error calling Gemini for IC Memo: {e}")
        return (
            f"### [오류] AI 분석 실행 중 예외 발생: {str(e)}\n\n"
            "#### 생성된 프롬프트 스크립트:\n\n"
            "```markdown\n" + prompt + "\n```"
        )
