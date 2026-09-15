import os
import sys
import sqlite3
from datetime import datetime

# pip install python-pptx
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
except ImportError:
    print("python-pptx 모듈이 설치되어 있지 않습니다. 아래 명령어로 설치해주세요.")
    print("pip install python-pptx")
    sys.exit(1)

def get_high_grade_news():
    db_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'dealsourcing.db')
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    
    # 등급 조건 완화: 최근 날짜, 높은 영향력 순으로 15개 가져오기
    query = """
    SELECT title, summary, link, pub_date, news_grade, impact_score, promising_industry
    FROM deal_articles
    ORDER BY pub_date DESC, impact_score DESC
    LIMIT 15
    """
    cursor.execute(query)
    results = cursor.fetchall()
    conn.close()
    
    return results

def add_header(slide, title_text, date_text):
    # 상단 다크 블루 배너
    banner = slide.shapes.add_shape(
        1, # msoShapeRectangle
        0, 0, Inches(13.333), Inches(1.2) # 16:9 widescreen width is 13.333 inches, height 7.5
    )
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(33, 47, 61) # Dark Blue/Slate
    banner.line.color.rgb = RGBColor(33, 47, 61)
    
    # 제목 텍스트 박스
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(10), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    # 날짜 텍스트 박스
    dateBox = slide.shapes.add_textbox(Inches(10.5), Inches(0.4), Inches(2.5), Inches(0.5))
    dtf = dateBox.text_frame
    dp = dtf.paragraphs[0]
    dp.text = f"조회일: {date_text}"
    dp.font.size = Pt(14)
    dp.font.color.rgb = RGBColor(200, 200, 200)
    dp.alignment = PP_ALIGN.RIGHT

def add_badge(slide, left, top, text, bg_color):
    badge = slide.shapes.add_shape(
        5, # msoShapeRoundedRectangle
        left, top, Inches(1.8), Inches(0.4)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = bg_color
    badge.line.color.rgb = bg_color
    
    tf = badge.text_frame
    tf.margin_bottom = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

def create_presentation(news_items):
    prs = Presentation()
    # 16:9 비율 설정
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_slide_layout = prs.slide_layouts[6] # Blank layout
    
    # 1. 제목 슬라이드
    title_slide = prs.slides.add_slide(blank_slide_layout)
    
    # 배경
    bg = title_slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(33, 47, 61)
    
    # 메인 타이틀
    txBox = title_slide.shapes.add_textbox(Inches(2), Inches(2.5), Inches(9.333), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "VC 딜소싱\n주요 유망 뉴스 요약 리포트"
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # 서브 타이틀
    subBox = title_slide.shapes.add_textbox(Inches(2), Inches(4.5), Inches(9.333), Inches(1))
    subtf = subBox.text_frame
    subp = subtf.paragraphs[0]
    subp.text = f"생성일: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    subp.font.size = Pt(20)
    subp.font.color.rgb = RGBColor(200, 200, 200)
    subp.alignment = PP_ALIGN.CENTER

    current_date_str = datetime.now().strftime('%Y-%m-%d')
    
    # 2. 내용 슬라이드
    for item in news_items:
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # item: title(0), summary(1), link(2), pub_date(3), news_grade(4), impact_score(5), industry(6)
        title_text = item[0][:50] + "..." if len(item[0]) > 50 else item[0]
        
        add_header(slide, title_text, current_date_str)
        
        # 메타데이터 뱃지 렌더링
        grade = item[4] if item[4] else "N/A"
        date_str = item[3] if item[3] else "날짜 미상"
        industry = item[6] if item[6] else "기타"
        
        add_badge(slide, Inches(0.5), Inches(1.5), f"일자: {date_str[:10]}", RGBColor(84, 153, 199))
        add_badge(slide, Inches(2.5), Inches(1.5), f"중요도: {grade}", RGBColor(231, 76, 60))
        add_badge(slide, Inches(4.5), Inches(1.5), f"산업: {industry}", RGBColor(39, 174, 96))
        
        # 요약 박스 배경
        sum_bg = slide.shapes.add_shape(5, Inches(0.5), Inches(2.2), Inches(12.333), Inches(3.8))
        sum_bg.fill.solid()
        sum_bg.fill.fore_color.rgb = RGBColor(242, 243, 244) # Light Gray
        sum_bg.line.color.rgb = RGBColor(208, 211, 212)
        
        # 요약 내용
        sumBox = slide.shapes.add_textbox(Inches(0.7), Inches(2.4), Inches(11.9), Inches(3.4))
        sumTf = sumBox.text_frame
        sumTf.word_wrap = True
        
        sum_title = sumTf.paragraphs[0]
        sum_title.text = "[핵심 요약]"
        sum_title.font.bold = True
        sum_title.font.size = Pt(20)
        sum_title.font.color.rgb = RGBColor(44, 62, 80)
        
        sum_content = sumTf.add_paragraph()
        sum_content.text = item[1] if item[1] else "요약 내용이 없습니다."
        sum_content.font.size = Pt(18)
        sum_content.font.color.rgb = RGBColor(52, 73, 94)
        sum_content.space_before = Pt(14)
        
        # 출처 링크
        linkBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(12.333), Inches(0.5))
        linkTf = linkBox.text_frame
        link_p = linkTf.paragraphs[0]
        link_p.text = f"출처: {item[2]}"
        link_p.font.size = Pt(14)
        link_p.font.color.rgb = RGBColor(41, 128, 185)

    output_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'reports')
    os.makedirs(output_dir, exist_ok=True)
    
    output_path = os.path.join(output_dir, f'news_briefing_{datetime.now().strftime("%Y%m%d_%H%M")}.pptx')
    
    prs.save(output_path)
    print(f"슬라이드 생성 완료: {output_path}")

if __name__ == "__main__":
    print("뉴스 데이터를 수집 중입니다...")
    news_data = get_high_grade_news()
    if not news_data:
        print("조건에 맞는 뉴스가 없습니다.")
    else:
        print(f"총 {len(news_data)}건의 뉴스를 바탕으로 슬라이드를 생성합니다.")
        create_presentation(news_data)
